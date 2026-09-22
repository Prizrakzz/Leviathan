"""Wave-1 regression: the SAGIS weekly-exports silver Batch task must build a LIVE S3 client on the
shadow/canonical publish path.

The ``s3_client=None`` defect (hard-wired into ``build_flat_publish`` calls, correct only for
dry-run) bit two producer families: a ``--publish-mode shadow`` run reached the publisher's staging
loop and crashed with ``'NoneType' object has no attribute 'put_object'``. The fix builds the client
conditionally (``None if dry-run else get_thread_local_s3_client``); this test locks ``main()``:

  * shadow builds EXACTLY one live client and stages under ``_shadow/`` (canonical untouched);
  * dry-run builds NO client (nothing written);
  * argparse accepts ``--publish-mode shadow``.

Pure/hermetic: an in-memory :class:`FakeS3` and a recording client factory replace all AWS. The
loader + transform are stubbed so no real bronze/S3/network is touched.
"""
from __future__ import annotations

import datetime as dt
import sys

import pandas as pd

from jobs.batch import sagis_weekly_exports_silver_task as task
from leviathan.storage.paths import silver_sagis_weekly_key
from tests.unit.silver.conftest import FakeS3

CANONICAL_KEY = silver_sagis_weekly_key("exports")


# --------------------------------------------------------------------------- canned silver frame
def _weekly_df() -> pd.DataFrame:
    """EXACTLY the ``silver_sagis_weekly_exports`` physical columns; the three value columns
    (prog_exports_mt, pct_of_prior_yr, z_vs_3yr_avg) are fully non-null so the 0.5 floor clears.
    """
    return pd.DataFrame({
        "season": ["2023-24", "2023-24"],
        "crop": ["maize", "maize"],
        "week_number": [10, 11],
        "week_ending": ["2023-12-08", "2023-12-15"],
        "week_ending_date": [dt.date(2023, 12, 8), dt.date(2023, 12, 15)],
        "prog_exports_mt": [125_000.0, 138_000.0],
        "pct_of_prior_yr": [104.5, 110.2],
        "z_vs_3yr_avg": [0.42, 0.61],
        "source": ["sagis_weekly", "sagis_weekly"],
    })


def _wire(monkeypatch):
    """Stub env + loader + transform + the S3 client factory. Returns ``(fake_s3, factory_calls)``."""
    monkeypatch.setattr(task, "load_env", lambda *a, **k: None)
    monkeypatch.setenv("AWS_REGION", "us-east-1")
    monkeypatch.setattr(task, "load_rows", lambda *a, **k: [])           # no S3 read in the loader
    monkeypatch.setattr(task, "transform_weekly_exports", lambda *a, **k: _weekly_df())
    s3 = FakeS3()
    calls: list[str] = []

    def _factory(region):
        calls.append(region)
        return s3

    # main() resolves ``get_thread_local_s3_client`` via a call-time import from leviathan.storage.s3,
    # so patch the source-module name (patching the task attribute would not take effect).
    monkeypatch.setattr("leviathan.storage.s3.get_thread_local_s3_client", _factory)
    return s3, calls


def test_shadow_builds_one_live_client_and_stages_to_shadow(monkeypatch):
    s3, calls = _wire(monkeypatch)
    monkeypatch.setattr(sys, "argv", ["prog", "--bucket", "leviathan-test",
                                      "--publish-mode", "shadow"])

    rc = task.main()

    assert rc == 0
    assert calls == ["us-east-1"]                     # exactly one live client, for the publish
    keys = s3.keys()
    assert any("_shadow" in k for k in keys)          # object staged shadow-first
    assert CANONICAL_KEY not in keys                  # canonical untouched in shadow mode


def test_dry_run_builds_no_client_and_writes_nothing(monkeypatch):
    s3, calls = _wire(monkeypatch)
    monkeypatch.setattr(sys, "argv", ["prog", "--bucket", "leviathan-test",
                                      "--publish-mode", "dry-run"])

    rc = task.main()

    assert rc == 0
    assert calls == []                                # dry-run stages nothing -> no client built
    assert s3.keys() == []                            # nothing written anywhere


def test_argparse_accepts_publish_mode_shadow(monkeypatch):
    monkeypatch.setattr(sys, "argv", ["prog", "--bucket", "leviathan-test",
                                      "--publish-mode", "shadow"])
    args = task._parse_args()
    assert args.publish_mode == "shadow"
    assert args.bucket == "leviathan-test"


# ------------------------------------------------------- RAW-workbook loader (census B9 / P9)
# THE PIN THIS SECTION REPLACES ITS PREDECESSOR WITH. Until 2026-09-22 these tests exercised a
# GOVERNED BRONZE schema, and they passed for months while the table was dead: nothing in the
# repository writes bronze/production/source=sagis_weekly/dataset=imp_exp_progressive/, so the
# producer re-derived 24 objects last written 2026-06-03 on every Friday fire, exited 0, and left a
# SERVED export-pace card whose newest week ended 2024-04-26 while RAW carried a workbook fetched
# 2026-09-18. A fixture of a schema no phase produces cannot catch that; a fixture of the RAW
# WORKBOOK the fetcher actually lands can.
def _write_sheet(ws, header, rows, title="WHEAT: RSA EXPORTS"):
    """One SAGIS export sheet: a title block, the header band, then the week rows."""
    ws.append([title])
    ws.append([])
    ws.append([None] + header)
    for r in rows:
        ws.append(r)


def _wheat_workbook() -> bytes:
    """Era A: ONE export sheet, the modern 'Progressive Total/Totaal' header."""
    import io as _io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RSA EXPORTS"
    _write_sheet(ws, ["Week", "BOTSWANA", "LESOTHO", "Week Total/Totaal",
                      "Progressive Total/Totaal"],
                 [[1, "27 Sep - 03 Oct/Okt 2025", 105, 857, 962, 962],
                  [2, "04 Oct/Okt - 10 Oct/Okt 2025", 140, 357, 497, 1459],
                  [3, "11 Oct/Okt - 17 Oct/Okt 2025", 36, 0, 36, 1495],
                  [None, "Total", 281, 1214, 1495, None]])
    # A re-export sheet and a harbour split, both of which would DOUBLE-COUNT if read.
    _write_sheet(wb.create_sheet("EXPORTS OF IMPORTED WHEAT"),
                 ["Week", "ZAMBIA", "Week Total/Totaal", "Progressive Total/Totaal"],
                 [[1, "27 Sep - 03 Oct/Okt 2025", 999, 999, 999],
                  [2, "04 Oct/Okt - 10 Oct/Okt 2025", 999, 999, 1998]])
    _write_sheet(wb.create_sheet("EXPORT PER HARBOUR"),
                 ["Week", "DURBAN", "Week Total/Totaal", "Progressive Total/Totaal"],
                 [[1, "27 Sep - 03 Oct/Okt 2025", 962, 962, 962],
                  [2, "04 Oct/Okt - 10 Oct/Okt 2025", 497, 497, 1459]])
    buf = _io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _maize_workbook(white_prog, yellow_prog) -> bytes:
    """Era B: the COLOUR-SPLIT maize workbook, and the exact shape of the live defect.

    2016-17 and 2019-20 canonical carried the WHITE sheet alone (32/32 and 50/50 rows equal to the
    white-only progressive total), so South Africa's yellow maize exports were absent from a series
    called 'maize export pace'. The header here prints 'Total/Totaal' TWICE -- week total, then
    progressive total -- which is the 2013-2018 era in which the column cannot be found by name."""
    import io as _io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "White RSA EXPORTS"
    _write_sheet(ws, ["Week", "BOTSWANA", "Total/Totaal", "Total/Totaal"],
                 [[i + 1, "week %d" % (i + 1), 10, 10, p] for i, p in enumerate(white_prog)],
                 title="WHITE MAIZE: RSA EXPORTS")
    _write_sheet(wb.create_sheet("YELLOW RSA EXPORTS"),
                 ["Week", "ZIMBABWE", "Total/Totaal", "Total/Totaal"],
                 [[i + 1, "week %d" % (i + 1), 10, 10, p] for i, p in enumerate(yellow_prog)],
                 title="YELLOW MAIZE: RSA EXPORTS")
    _write_sheet(wb.create_sheet("White IMPORTS FOR RSA"),
                 ["Week", "MEXICO", "Total/Totaal", "Total/Totaal"],
                 [[i + 1, "week %d" % (i + 1), 0, 0, 500_000] for i in range(len(white_prog))],
                 title="WHITE MAIZE: IMPORTS")
    buf = _io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_the_dead_bronze_prefix_is_gone_from_the_producer():
    """The recurrence pin: this producer may never again read a prefix no phase writes."""
    from pathlib import Path
    src = Path(task.__file__).read_text(encoding="utf-8")
    code = "\n".join(line for line in src.splitlines() if not line.lstrip().startswith("#"))
    assert '_BRONZE_PREFIX' not in code
    assert task._RAW_PREFIX == "raw/production/source=sagis_weekly/dataset=imp_exp_progressive/"


def test_reads_the_progressive_total_and_ignores_reexport_and_harbour_sheets():
    rows = task.read_workbook_exports(_wheat_workbook(),
                                      "IMP-EXP_Progressive_Koring_2025-2026_3.xlsx")
    assert [(r["week_number"], r["prog_exports_mt"]) for r in rows] == [
        (1, 962.0), (2, 1459.0), (3, 1495.0)]
    # the footer 'Total' row carries no week number, so it never becomes an observation
    assert all(r["week_number"] <= 3 for r in rows)
    assert rows[0]["week_ending"] == "27 Sep - 03 Oct/Okt 2025"


def test_colour_split_maize_sheets_are_summed_not_taken_white_only():
    """B9's measured defect, as a fixture: white alone is 100, the national total is 175."""
    rows = task.read_workbook_exports(_maize_workbook([100.0, 250.0], [75.0, 190.0]),
                                      "IMP-EXP_Progressive_Mielies_2019-2020_2.xlsx")
    assert [(r["week_number"], r["prog_exports_mt"]) for r in rows] == [(1, 175.0), (2, 440.0)]


def test_load_rows_keeps_both_crops_of_one_season_and_takes_the_widest_snapshot(monkeypatch):
    """Two facts in one run, because they fail together.

    (1) The authoritative snapshot is per (season, crop): the two-week maize file must lose to the
        three-week one. (2) Wheat must SURVIVE the maize selection --
        ``sagis_weekly_exports.select_authoritative_snapshot`` keys its winner on ``season`` ALONE,
        so a per-season selection with real file identities drops one crop of every season
        outright. That is why the selection lives in ``load_rows``."""
    import datetime as _dt
    import io as _io

    prefix = task._RAW_PREFIX
    bodies = {
        prefix + "crop=maize/IMP-EXP_Progressive_Mielies_2025-2026_2.xlsx":
            _maize_workbook([1.0, 2.0], [1.0, 2.0]),
        prefix + "crop=maize/IMP-EXP_Progressive_Mielies_2025-2026_3.xlsx":
            _maize_workbook([10.0, 20.0, 30.0], [1.0, 2.0, 3.0]),
        prefix + "crop=wheat/IMP-EXP_Progressive_Koring_2025-2026_3.xlsx": _wheat_workbook(),
    }
    stamps = {k: _dt.datetime(2026, 9, 1 + i, tzinfo=_dt.timezone.utc)
              for i, k in enumerate(sorted(bodies))}

    class _S3:
        def get_object(self, Bucket, Key):
            return {"Body": _io.BytesIO(bodies[Key])}

    monkeypatch.setattr("leviathan.storage.s3.get_thread_local_s3_client", lambda r: _S3())
    monkeypatch.setattr("leviathan.storage.s3.list_s3_keys_with_mtime", lambda *a, **k: stamps)

    rows = task.load_rows("bucket", "us-east-1")
    by_crop: dict[str, list] = {}
    for r in rows:
        by_crop.setdefault(r.crop, []).append(r)
    assert set(by_crop) == {"maize", "wheat"}, "a per-SEASON selection drops one crop"
    assert [r.prog_exports_mt for r in sorted(by_crop["maize"], key=lambda r: r.week_number)] == \
        [11.0, 22.0, 33.0]                     # the WIDEST snapshot won, and both colours summed
    assert [r.prog_exports_mt for r in sorted(by_crop["wheat"], key=lambda r: r.week_number)] == \
        [962.0, 1459.0, 1495.0]
    assert {r.season for r in rows} == {"2025-26"}
    assert all(r.is_total for r in rows)
